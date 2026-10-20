#!/usr/bin/env python3
"""
Apply 50% Numerical Questions to All 21 HSC Physics Workbooks.
- Rows 2 to 51: Retained as theoretical / conceptual proofs.
- Rows 52 to 101: 50 Numerical questions with direct answers, step-by-step explanations, and vector diagram delimiters.
- Column 8: Primary Image URL (##PRESET:...[size:medium]## or ##GEO:...[size:medium]##)
- Column 9: Teacher Note / Detailed Explanation
- Column 10: Direct Model Answer
"""

import sys
import os
import openpyxl

sys.path.append('/Users/md.rofazhasanrafiu/coding/digital_school/scripts')

from numericals_ch01 import get_chapter_01_numericals
from numericals_ch02_ch03 import get_chapter_02_numericals, get_chapter_03_numericals
from numericals_ch04_ch05 import get_chapter_04_numericals, get_chapter_05_numericals
from numericals_ch06_ch07 import get_chapter_06_numericals, get_chapter_07_numericals
from numericals_ch08_ch10 import get_chapter_08_numericals, get_chapter_09_numericals, get_chapter_10_numericals
from numericals_2nd_ch01_ch03 import get_chapter_01_2nd_numericals, get_chapter_02_2nd_numericals, get_chapter_03_2nd_numericals
from numericals_2nd_ch04_ch06 import get_chapter_04_2nd_numericals, get_chapter_05_2nd_numericals, get_chapter_06_2nd_numericals
from numericals_2nd_ch07_ch09 import get_chapter_07_2nd_numericals, get_chapter_08_2nd_numericals, get_chapter_09_2nd_numericals
from numericals_2nd_ch10_ch11 import get_chapter_10_2nd_numericals, get_chapter_11_2nd_numericals

BASE_DIR = '/Users/md.rofazhasanrafiu/coding'

CHAPTER_CONFIG = [
    # 1st Paper
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_01_Physical_World_and_Measurement.xlsx',
        'generator': get_chapter_01_numericals,
        'name': '1st Paper Chapter 01 (Physical World and Measurement)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_02_Vector.xlsx',
        'generator': get_chapter_02_numericals,
        'name': '1st Paper Chapter 02 (Vector)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_03_Dynamics.xlsx',
        'generator': get_chapter_03_numericals,
        'name': '1st Paper Chapter 03 (Dynamics)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_04_Newtonian_Mechanics.xlsx',
        'generator': get_chapter_04_numericals,
        'name': '1st Paper Chapter 04 (Newtonian Mechanics)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_05_Work_Energy_and_Power.xlsx',
        'generator': get_chapter_05_numericals,
        'name': '1st Paper Chapter 05 (Work, Energy and Power)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_06_Gravitation_and_Gravity.xlsx',
        'generator': get_chapter_06_numericals,
        'name': '1st Paper Chapter 06 (Gravitation and Gravity)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_07_Structural_Properties_of_Matter.xlsx',
        'generator': get_chapter_07_numericals,
        'name': '1st Paper Chapter 07 (Structural Properties of Matter)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_08_Periodic_Motion.xlsx',
        'generator': get_chapter_08_numericals,
        'name': '1st Paper Chapter 08 (Periodic Motion)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_09_Waves.xlsx',
        'generator': get_chapter_09_numericals,
        'name': '1st Paper Chapter 09 (Waves)'
    },
    {
        'file': 'HSC_Physics_1st_Paper_Chapter_10_Ideal_Gas_and_Kinetic_Theory_of_Gases.xlsx',
        'generator': get_chapter_10_numericals,
        'name': '1st Paper Chapter 10 (Ideal Gas and Kinetic Theory of Gases)'
    },
    # 2nd Paper
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_01_Thermodynamics.xlsx',
        'generator': get_chapter_01_2nd_numericals,
        'name': '2nd Paper Chapter 01 (Thermodynamics)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_02_Static_Electricity.xlsx',
        'generator': get_chapter_02_2nd_numericals,
        'name': '2nd Paper Chapter 02 (Static Electricity)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_03_Current_Electricity.xlsx',
        'generator': get_chapter_03_2nd_numericals,
        'name': '2nd Paper Chapter 03 (Current Electricity)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_04_Magnetic_Effects_of_Current_and_Magnetism.xlsx',
        'generator': get_chapter_04_2nd_numericals,
        'name': '2nd Paper Chapter 04 (Magnetic Effects of Current & Magnetism)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_05_Electromagnetic_Induction_and_Alternating_Current.xlsx',
        'generator': get_chapter_05_2nd_numericals,
        'name': '2nd Paper Chapter 05 (Electromagnetic Induction & AC)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_06_Geometrical_Optics.xlsx',
        'generator': get_chapter_06_2nd_numericals,
        'name': '2nd Paper Chapter 06 (Geometrical Optics)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_07_Physical_Optics.xlsx',
        'generator': get_chapter_07_2nd_numericals,
        'name': '2nd Paper Chapter 07 (Physical Optics)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_08_Introduction_to_Modern_Physics.xlsx',
        'generator': get_chapter_08_2nd_numericals,
        'name': '2nd Paper Chapter 08 (Introduction to Modern Physics)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_09_Atomic_Model_and_Nuclear_Physics.xlsx',
        'generator': get_chapter_09_2nd_numericals,
        'name': '2nd Paper Chapter 09 (Atomic Model & Nuclear Physics)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_10_Semiconductor_and_Electronics.xlsx',
        'generator': get_chapter_10_2nd_numericals,
        'name': '2nd Paper Chapter 10 (Semiconductor & Electronics)'
    },
    {
        'file': 'HSC_Physics_2nd_Paper_Chapter_11_Astronomy.xlsx',
        'generator': get_chapter_11_2nd_numericals,
        'name': '2nd Paper Chapter 11 (Astronomy)'
    }
]

def process_chapter(config):
    filepath = os.path.join(BASE_DIR, config['file'])
    if not os.path.exists(filepath):
        print(f"[ERROR] File not found: {filepath}")
        return False

    numericals = config['generator']()
    if len(numericals) != 50:
        print(f"[ERROR] Expected 50 numerical items, got {len(numericals)} for {config['name']}")
        return False

    wb = openpyxl.load_workbook(filepath)
    ws = wb['Template'] if 'Template' in wb.sheetnames else wb.active

    # Extract default class_name and subject from row 2
    class_name = ws.cell(row=2, column=2).value or "HSC / Alim"
    subject_name = ws.cell(row=2, column=3).value or "Physics"

    # Inject into rows 52 to 101
    for idx, item in enumerate(numericals):
        row_num = 52 + idx # 52 to 101 inclusive

        ws.cell(row=row_num, column=1, value='SQ')
        ws.cell(row=row_num, column=2, value=class_name)
        ws.cell(row=row_num, column=3, value=subject_name)
        ws.cell(row=row_num, column=4, value=item['topic'])
        ws.cell(row=row_num, column=5, value=item['difficulty'])
        ws.cell(row=row_num, column=6, value=item['marks'])
        ws.cell(row=row_num, column=7, value=item['questionText'])
        ws.cell(row=row_num, column=8, value=item.get('primaryImage') or '')
        ws.cell(row=row_num, column=9, value=item['explanation'])
        ws.cell(row=row_num, column=10, value=item['modelAnswer'])

        # Clear MCQ options if any existed in these rows
        for c in range(11, 15):
            ws.cell(row=row_num, column=c, value='')

    wb.save(filepath)
    print(f"[SUCCESS] Updated {config['name']} -> 50 numericals injected (Rows 52-101), total rows = {ws.max_row}")
    return True

def main():
    print("=" * 60)
    print("Starting injection of 50% numerical questions across all 21 workbooks...")
    print("=" * 60)

    success_count = 0
    for cfg in CHAPTER_CONFIG:
        if process_chapter(cfg):
            success_count += 1

    print("=" * 60)
    print(f"Injection Complete: {success_count}/{len(CHAPTER_CONFIG)} workbooks updated successfully.")
    print("=" * 60)

if __name__ == '__main__':
    main()
