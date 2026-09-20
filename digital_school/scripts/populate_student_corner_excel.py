import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
from datetime import datetime

src_path = '/Users/md.rofazhasanrafiu/coding/Unified_Admission_Master_Tracker_2026.xlsx'
dest_path = '/Users/md.rofazhasanrafiu/coding/student_corner_template.xlsx'

src_wb = openpyxl.load_workbook(src_path, data_only=True)
ws_105 = src_wb['⚡ 105D Revision Challenge']

rows_to_import = []

for r in range(10, 115):
    sprint_day = ws_105.cell(r, 1).value
    dt = ws_105.cell(r, 2).value
    day_name = ws_105.cell(r, 3).value
    phase = str(ws_105.cell(r, 4).value or '').strip()
    task1 = str(ws_105.cell(r, 5).value or '').strip()
    task2 = str(ws_105.cell(r, 6).value or '').strip()
    task3 = str(ws_105.cell(r, 7).value or '').strip()
    task4 = str(ws_105.cell(r, 8).value or '').strip()
    
    if not dt:
        continue
        
    date_str = dt.strftime('%Y-%m-%d') if hasattr(dt, 'strftime') else str(dt)[:10]
    
    # --- Task 1: Morning Batch / Syllabus Study (STUDY) ---
    sub_matches = re.findall(r'•\s*([^:]+):\s*([^•\n]+)', task1)
    if sub_matches:
        subjects_list = [m[0].strip() for m in sub_matches]
        topics_list = [re.sub(r'\[Done\]', '', m[1]).strip() for m in sub_matches]
        t1_subject = ' & '.join(subjects_list)
        t1_topics = '; '.join(topics_list)
        t1_title = f'Batch Study: {t1_subject}'
    else:
        clean_t1 = re.sub(r'\[Done\]', '', task1).strip()
        t1_title = 'Concept Mastery & Practice Drill' if 'Concept Mastery' in clean_t1 else (clean_t1[:45] if clean_t1 else 'Morning Batch & Syllabus Study')
        t1_subject = 'Physics, Chemistry & Math'
        t1_topics = clean_t1
        
    rows_to_import.append({
        'date': date_str,
        'title': t1_title,
        'category': 'STUDY',
        'subject': t1_subject,
        'topics': t1_topics,
        'duration': 180,
        'time': '08:30',
        'priority': 'HIGH',
        'notes': f'{sprint_day} ({day_name}) • {phase} • Morning batch syllabus study deadline: 11:30'
    })
    
    # --- Task 2: Spaced Active Recall (REVISION) ---
    is_pre_exam = 'প্রাক-পরীক্ষা' in task2 or 'প্রি-এক্সাম' in task2
    clean_t2 = re.sub(r'^[⚡📐⚗️♾️🧬🧪🏆\s]+', '', task2).strip()
    
    if is_pre_exam:
        t2_title = 'Pre-Exam Formula Scan & Readiness'
        t2_duration = 30
        t2_priority = 'URGENT'
        t2_subject = 'Pre-Exam Quick Revision'
    else:
        box_m = re.search(r'\((বক্স\s*\d+)\)', task2)
        box_str = f' [{box_m.group(1)}]' if box_m else ''
        t2_title = f'Spaced Active Recall{box_str}'
        t2_duration = 60
        t2_priority = 'HIGH'
        t2_subject = 'Spaced Repetition & Formula Retrieval'
        
    # Topic cleanup
    topic_summary = re.sub(r'^(রিকল\s*\([^)]+\):|প্রাক-পরীক্ষা\s*[^:]+:|প্রি-এক্সাম\s*[^:]+:)', '', clean_t2).strip()
    rows_to_import.append({
        'date': date_str,
        'title': t2_title,
        'category': 'REVISION',
        'subject': t2_subject,
        'topics': topic_summary[:120],
        'duration': t2_duration,
        'time': '14:30',
        'priority': t2_priority,
        'notes': f'{sprint_day} ({day_name}) • Pen-paper active retrieval drill deadline: 16:30'
    })
    
    # --- Task 3: Exam Simulation or Question Bank (EXAM_PREP) ---
    if 'Engineering 230M' in task3:
        code_m = re.search(r'Engineering\s*(W-\d+)', task3)
        code_str = f' ({code_m.group(1)})' if code_m else ''
        t3_title = f'Engineering 230M Exam Simulation{code_str}'
        t3_subject = 'Engineering Admission (BUET/CKRUET)'
        t3_duration = 150
        t3_priority = 'URGENT'
        clean_topics = re.sub(r'⚙️\s*Engineering 230M:\s*★\s*[^\n]+\n', '', task3)
        clean_topics = re.sub(r'•\s*', '', clean_topics).replace('\n', '; ').strip()
    elif 'Varsity 120M' in task3:
        code_m = re.search(r'Varsity\s*(W-\d+)', task3)
        code_str = f' ({code_m.group(1)})' if code_m else ''
        t3_title = f'Varsity 120M Exam Simulation{code_str}'
        t3_subject = 'Varsity \'Ka\' Admission (DU/RU/JU)'
        t3_duration = 90
        t3_priority = 'URGENT'
        clean_topics = '6 Subjects: PCM + Bio + Eng + ICT; Triad: 37 Easy : 60 Med : 23 Tough'
    else:
        t3_title = 'BUET & DU Past Questions Speed Run (30 Qs)'
        t3_subject = 'BUET & DU Question Bank'
        t3_duration = 45
        t3_priority = 'HIGH'
        clean_topics = 'Solve 30 BUET & DU Past Questions under Strict 45-Min Timer'
        
    rows_to_import.append({
        'date': date_str,
        'title': t3_title,
        'category': 'EXAM_PREP',
        'subject': t3_subject,
        'topics': clean_topics[:140],
        'duration': t3_duration,
        'time': '17:00',
        'priority': t3_priority,
        'notes': f'{sprint_day} ({day_name}) • High-stakes timed simulation cutoff: 19:30'
    })
    
    # --- Task 4: 5 Daily Prayers, Brain Nutrition & Curfew (SALAT) ---
    rows_to_import.append({
        'date': date_str,
        'title': '5 Waqt Salat, Nutrition & Night Curfew',
        'category': 'SALAT',
        'subject': 'Islamic Routine & Wellness',
        'topics': '5 Waqt Salat in Mosque; 3L Hydration; 0 Social Media; 23:00 Sleep Lockout',
        'duration': 45,
        'time': '20:00',
        'priority': 'HIGH',
        'notes': f'{sprint_day} ({day_name}) • Circadian sleep lockout & digital blackout by 23:00'
    })

print(f'Total Generated Challenges: {len(rows_to_import)} rows')

# Now create target workbook
wb_out = openpyxl.Workbook()

# Sheet 1: Daily Challenges
ws_out = wb_out.active
ws_out.title = 'Daily Challenges'

headers = [
    'Date (YYYY-MM-DD)',
    'Challenge Title',
    'Category',
    'Subject',
    'Topics (Semicolon separated)',
    'Duration Minutes',
    'Scheduled Time (HH:mm)',
    'Priority',
    'Notes'
]

ws_out.append(headers)

# Styling header
header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

for col_num in range(1, len(headers) + 1):
    cell = ws_out.cell(1, col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align

ws_out.row_dimensions[1].height = 28

# Add rows
category_colors = {
    'STUDY': 'EFF6FF',      # light blue
    'REVISION': 'FDF4FF',   # light purple
    'EXAM_PREP': 'FFF1F2',  # light rose
    'SALAT': 'F0FDF4',      # light green
}

for r_idx, item in enumerate(rows_to_import, start=2):
    ws_out.append([
        item['date'],
        item['title'],
        item['category'],
        item['subject'],
        item['topics'],
        item['duration'],
        item['time'],
        item['priority'],
        item['notes']
    ])
    
    ws_out.row_dimensions[r_idx].height = 22
    row_bg = category_colors.get(item['category'], 'FFFFFF')
    row_fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type='solid')
    
    for c_idx in range(1, len(headers) + 1):
        cell = ws_out.cell(r_idx, c_idx)
        cell.font = Font(name='Arial', size=10)
        cell.border = thin_border
        cell.fill = row_fill
        
        # Center align date, duration, time, priority, category
        if c_idx in [1, 3, 6, 7, 8]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Adjust column widths
col_widths = [18, 36, 16, 28, 45, 18, 22, 16, 45]
for i, width in enumerate(col_widths, start=1):
    col_letter = get_column_letter(i)
    ws_out.column_dimensions[col_letter].width = width

# Sheet 2: Allowed Values Reference
ws_ref = wb_out.create_sheet('Allowed Values Reference')
ws_ref.append(['Allowed Categories', 'Allowed Priorities'])

ref_header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
for col_num in [1, 2]:
    cell = ws_ref.cell(1, col_num)
    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    cell.fill = ref_header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')

categories = ['STUDY', 'CODING', 'SALAT', 'QURAN', 'HABIT', 'EXERCISE', 'READING', 'DIET', 'ASSIGNMENT', 'REVISION', 'EXAM_PREP', 'CUSTOM']
priorities = ['LOW', 'MEDIUM', 'HIGH', 'URGENT']

max_l = max(len(categories), len(priorities))
for i in range(max_l):
    cat = categories[i] if i < len(categories) else ''
    pri = priorities[i] if i < len(priorities) else ''
    ws_ref.append([cat, pri])
    r_num = i + 2
    ws_ref.cell(r_num, 1).alignment = Alignment(horizontal='center')
    ws_ref.cell(r_num, 2).alignment = Alignment(horizontal='center')
    ws_ref.cell(r_num, 1).border = thin_border
    ws_ref.cell(r_num, 2).border = thin_border

ws_ref.column_dimensions['A'].width = 25
ws_ref.column_dimensions['B'].width = 25

wb_out.save(dest_path)
print(f'Successfully populated {dest_path} with {len(rows_to_import)} challenge rows across 105 admission days!')
