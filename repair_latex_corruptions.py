#!/usr/bin/env python3
"""
Repairs all LaTeX syntax corruptions caused by unescaped Python control sequences across all 21 workbooks.
Fixes:
- \x0c + rac{ -> \frac{
- \x07 + lpha -> \alpha
- \x08 + eta -> \beta
- \x0b + ec -> \vec
- \t + heta -> \theta
- \t + ext -> \text
- \t + au -> \tau
- \t + imes -> \times
- \t + an -> \tan
- Raw 'rac{' -> '\frac{'
- Raw 'lpha' -> '\alpha'
"""

import os
import glob
import re
import openpyxl

FOLDER = '/Users/md.rofazhasanrafiu/coding/untitled folder'

def repair_text(text: str) -> str:
    if not text:
        return text
    
    # Replace control characters
    text = text.replace('\x0c', r'\f')
    text = text.replace('\x07', r'\a')
    text = text.replace('\x08', r'\b')
    text = text.replace('\x0b', r'\v')
    
    # Fix tab-corrupted LaTeX tokens
    tab_tokens = {
        '\theta': r'\theta',
        '\text': r'\text',
        '\tau': r'\tau',
        '\times': r'\times',
        '\tan': r'\tan',
        '\triangle': r'\triangle',
        '\two': r'\two',
        '\to': r'\to',
    }
    # Direct string replacement of tab + substring
    text = text.replace('\theta', r'\theta')
    text = text.replace('\text', r'\text')
    text = text.replace('\tau', r'\tau')
    text = text.replace('\times', r'\times')
    text = text.replace('\tan', r'\tan')
    text = text.replace('\to', r'\to')
    
    # Fix broken tokens where backslash was stripped
    text = re.sub(r'(?<![a-zA-Z\\])rac\{', r'\\frac{', text)
    text = re.sub(r'(?<![a-zA-Z\\])lpha(?![a-zA-Z])', r'\\alpha', text)
    text = re.sub(r'(?<![a-zA-Z\\])eta(?=[_ \^,\.\$\)])', r'\\beta', text)
    text = re.sub(r'(?<![a-zA-Z\\])heta(?=[_ \^,\.\$\)])', r'\\theta', text)
    
    # Clean any unintended double-escapes like \\frac -> \frac
    for kw in ['frac', 'alpha', 'beta', 'vec', 'theta', 'tau', 'text', 'times', 'tan', 'to']:
        text = text.replace('\\\\' + kw, '\\' + kw)
        
    return text

def main():
    files = sorted(glob.glob(os.path.join(FOLDER, '*.xlsx')))
    files = [f for f in files if not os.path.basename(f).startswith('~$')]
    
    total_repaired = 0
    for f in files:
        wb = openpyxl.load_workbook(f)
        ws = wb['Template'] if 'Template' in wb.sheetnames else wb.active
        modified = False
        
        for r in range(2, 102):
            for c in [7, 8, 9, 10]:
                orig = ws.cell(r, c).value
                if isinstance(orig, str):
                    repaired = repair_text(orig)
                    if repaired != orig:
                        ws.cell(r, c, value=repaired)
                        modified = True
                        total_repaired += 1
                        
        if modified:
            wb.save(f)
            print(f"Repaired: {os.path.basename(f)}")
            
    print(f"\nAll files processed! Total cell values repaired: {total_repaired}")

if __name__ == '__main__':
    main()
